import io
import json
import hmac
import hashlib
import re
from collections import Counter
from datetime import date, datetime, timedelta
from zoneinfo import ZoneInfo

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
import psycopg


# =========================
# 基本設定
# =========================
st.set_page_config(
    page_title="超慧科技｜生管資材戰情看板",
    page_icon="📊",
    layout="wide",
)

TAIPEI = ZoneInfo("Asia/Taipei")

MODULES = {
    "production": "2026排程",
    "material": "缺料追蹤",
    "iqc": "IQC 待驗",
    "warehouse": "倉庫概況",
}

FIELDS = [
    "編號", "專案", "料號或機型", "名稱", "供應商或地點",
    "數量", "完成數量", "需求日期", "預計日期",
    "狀態", "優先級", "負責人", "備註"
]

STATES = {
    "production": ["未開始", "組立中", "檢驗中", "已完成", "暫停"],
    "material": ["待確認", "齊料", "待回覆", "採購中", "已到貨", "已結案"],
    "iqc": ["待驗", "檢驗中", "合格", "不合格", "已結案"],
    "warehouse": ["使用中", "空儲位", "停用"],
}

DONE = {
    "production": {"已完成"},
    "material": {"齊料", "已到貨", "已結案"},
    "iqc": {"合格", "已結案"},
    "warehouse": {"空儲位", "停用"},
}

MATERIAL_COMPARE_FIELDS = [
    "Frame/ Frame set", "PU", "Facility", "其他託外模組",
    "RB", "LP", "AL", "FFU", "X-table", "加工件", "市購件", "齊料"
]

# Excel 匯入時保留人工操作記憶，不被新 Excel 蓋掉
MANUAL_MEMORY_FIELDS = ["優先級", "備註"]
MANUAL_MEMORY_INTERNAL = [
    "_manual_note_updated_at",
    "_last_imported_at",
    "_progress_changed_at",
    "_previous_progress",
    "_last_compare_changed",
    "_progress_diff",
    "_material_changed_at",
    "_material_diff",
]


# =========================
# 共用工具
# =========================
def now_dt():
    return datetime.now(TAIPEI)

def now_stamp():
    return now_dt().isoformat(timespec="seconds")

def today():
    return now_dt().date()

def txt(v):
    if v is None:
        return ""
    if isinstance(v, (date, datetime)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()

def parse_day(v):
    s = str(v or "").strip()
    if not s:
        return None
    try:
        return date.fromisoformat(s[:10])
    except Exception:
        return None

def parse_excel_date(v):
    s = txt(v)
    if s in ["", "-", "--", "TBD", "託"]:
        return "", ""
    for fmt in ["%Y-%m-%d", "%Y/%m/%d", "%Y%m%d", "%Y%m/%d", "%Y/%m%d"]:
        try:
            d = datetime.strptime(s.splitlines()[-1].strip(), fmt).date().isoformat()
            return d, ("多行日期採最後一行" if "\n" in s else "")
        except ValueError:
            pass
    return "", "日期無法辨識，請確認原文"

def elapsed_days(stamp):
    s = str(stamp or "").strip()
    if not s:
        return 0
    try:
        d = datetime.fromisoformat(s).date()
    except Exception:
        try:
            d = date.fromisoformat(s[:10])
        except Exception:
            return 0
    return max((today() - d).days, 0)

def display_value(v):
    s = str(v or "").strip()
    return s if s else "空白"

def validate_row(row, module):
    r = dict(row)
    for f in FIELDS:
        if f not in r:
            r[f] = ""
    if not str(r.get("編號", "")).strip():
        raise ValueError("編號不可空白")

    for f in ["數量", "完成數量"]:
        try:
            n = float(r.get(f) or 0)
        except Exception:
            raise ValueError(f"{f} 須為數字")
        if n < 0:
            raise ValueError(f"{f} 不可小於 0")
        r[f] = n

    if r["完成數量"] > r["數量"]:
        raise ValueError("完成數量不可大於數量")

    for f in ["需求日期", "預計日期"]:
        if r.get(f):
            d = parse_day(r.get(f))
            if not d:
                raise ValueError(f"{f} 須為 YYYY-MM-DD")
            r[f] = d.isoformat()

    if r.get("狀態") not in STATES[module]:
        raise ValueError("狀態不符此模組")
    if r.get("優先級") not in ["一般", "急件"]:
        raise ValueError("優先級須為一般或急件")

    for f in FIELDS:
        if len(str(r.get(f, ""))) > 2000:
            raise ValueError(f"{f} 內容過長")

    return r


# =========================
# PostgreSQL 持久記憶
# Excel 是主資料來源；DB 只負責保存操作記憶與比較基準
# =========================
def db_connect(url):
    return psycopg.connect(url, connect_timeout=15)

def db_initialize(url):
    with db_connect(url) as c:
        c.execute("""
            CREATE TABLE IF NOT EXISTS dashboard_records(
                module TEXT NOT NULL,
                id TEXT NOT NULL,
                payload TEXT NOT NULL,
                version INTEGER NOT NULL DEFAULT 1,
                PRIMARY KEY(module, id)
            )
        """)
        # 舊版本資料庫可直接升級，不清空舊資料
        c.execute("""
            ALTER TABLE dashboard_records
            ADD COLUMN IF NOT EXISTS updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
        """)
        c.execute("""
            CREATE TABLE IF NOT EXISTS dashboard_meta(
                meta_key TEXT PRIMARY KEY,
                meta_value TEXT NOT NULL,
                updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
            )
        """)

def db_load(url):
    data = {m: {} for m in MODULES}
    with db_connect(url) as c:
        rows = c.execute("""
            SELECT module, id, payload, version
            FROM dashboard_records
        """).fetchall()
    for module, rid, payload, version in rows:
        if module in data:
            try:
                data[module][rid] = (json.loads(payload), int(version))
            except Exception:
                pass
    return data

def db_set_meta(url, key, value):
    text_value = json.dumps(value, ensure_ascii=False) if not isinstance(value, str) else value
    with db_connect(url) as c:
        c.execute("""
            INSERT INTO dashboard_meta(meta_key, meta_value, updated_at)
            VALUES(%s, %s, NOW())
            ON CONFLICT(meta_key)
            DO UPDATE SET meta_value=EXCLUDED.meta_value, updated_at=NOW()
        """, (key, text_value))

def db_get_all_meta(url):
    result = {}
    with db_connect(url) as c:
        rows = c.execute("""
            SELECT meta_key, meta_value, updated_at
            FROM dashboard_meta
            ORDER BY updated_at DESC
        """).fetchall()
    for k, v, ts in rows:
        try:
            parsed = json.loads(v)
        except Exception:
            parsed = v
        result[k] = {"value": parsed, "updated_at": str(ts)}
    return result

def db_save_manual(url, module, row, expected_version=None):
    row = validate_row(row, module)
    rid = row["編號"]
    payload = json.dumps(row, ensure_ascii=False)

    with db_connect(url) as c:
        if expected_version is None:
            cur = c.execute("""
                INSERT INTO dashboard_records(module, id, payload, version, updated_at)
                VALUES(%s, %s, %s, 1, NOW())
                ON CONFLICT(module, id) DO NOTHING
            """, (module, rid, payload))
            if cur.rowcount != 1:
                raise ValueError("編號已存在，請重新整理後再操作。")
        else:
            cur = c.execute("""
                UPDATE dashboard_records
                SET payload=%s, version=version+1, updated_at=NOW()
                WHERE module=%s AND id=%s AND version=%s
            """, (payload, module, rid, int(expected_version)))
            if cur.rowcount != 1:
                raise ValueError("資料已被其他使用者修改，請重新整理後再操作。")

def merge_manual_memory(new_row, old_row):
    if not old_row:
        return new_row
    merged = dict(new_row)

    # 人工操作記憶優先保留
    for f in MANUAL_MEMORY_FIELDS:
        old_val = old_row.get(f)
        if old_val not in [None, ""]:
            merged[f] = old_val

    # 內部比較記憶讓比較函式自行更新，其餘舊資訊保留
    for k, v in old_row.items():
        if k.startswith("_") and k not in merged:
            merged[k] = v

    return merged

def db_save_import(url, module, imported_rows, existing_entries):
    if module not in ["production", "material"]:
        raise ValueError("此模組不支援 Excel 匯入")

    saved = 0
    stamp = now_stamp()

    with db_connect(url) as c:
        for incoming in imported_rows:
            rid = incoming["編號"]
            old_pair = existing_entries.get(rid)
            old = old_pair[0] if old_pair else None

            row = merge_manual_memory(incoming, old)

            if module == "production":
                row = compare_schedule_row(row, old, stamp)
            else:
                row = compare_material_row(row, old, stamp)

            row = validate_row(row, module)
            payload = json.dumps(row, ensure_ascii=False)

            c.execute("""
                INSERT INTO dashboard_records(module, id, payload, version, updated_at)
                VALUES(%s, %s, %s, 1, NOW())
                ON CONFLICT(module, id)
                DO UPDATE SET
                    payload=EXCLUDED.payload,
                    version=dashboard_records.version+1,
                    updated_at=NOW()
            """, (module, rid, payload))
            saved += 1

    db_set_meta(url, f"last_import_{module}", {
        "time": stamp,
        "count": saved,
        "source": "Excel",
        "note": "Excel為主資料來源；舊資料未包含的項目不刪除；人工備註/優先級保留。",
    })
    return saved


# =========================
# Excel 解析
# =========================
SCHEDULE_FIELDS = [
    "製令","客戶","P/N","Type","Category","組立地點","組立人員",
    "組立進度","備註","發料日","入庫日","保稅","客戶入庫日"
]

def read_schedule(source):
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if "2026排程" not in wb.sheetnames:
            raise ValueError("找不到「2026排程」工作表")

        ws = wb["2026排程"]
        if ws.max_row > 30000:
            raise ValueError("超過 30000 列，請縮小資料範圍")

        iterator = ws.iter_rows(max_col=15, values_only=True)
        header = [re.sub(r"\s+", "", txt(v)) for v in next(iterator)]
        expected = [re.sub(r"\s+", "", x) for x in SCHEDULE_FIELDS]

        if header[:13] != expected:
            raise ValueError("前 13 欄與 2026排程格式不符，請保留原欄位順序")

        records, warnings = [], []
        occ = Counter()

        for line, values in enumerate(iterator, 2):
            if not txt(values[0]):
                continue

            raw = dict(zip(SCHEDULE_FIELDS, map(txt, values[:13])))
            r = dict(raw)
            notes = []

            key = json.dumps(
                [raw[k] for k in ["製令", "P/N", "Type", "Category"]],
                ensure_ascii=False
            )
            occ[key] += 1
            uid = "SCH26-" + hashlib.sha256(
                (key + "#" + str(occ[key])).encode("utf-8")
            ).hexdigest()[:24]

            for f in ["發料日", "入庫日", "客戶入庫日"]:
                r[f], note = parse_excel_date(raw[f])
                if note:
                    notes.append(f"{f}：{note}")

            progress = str(r.get("組立進度", "")).strip()
            if progress in ["已完工", "已完成"]:
                state = "已完成"
            elif "暫停" in progress:
                state = "暫停"
            elif progress in ["Q", "待Q", "R", "Q / R"]:
                state = "檢驗中"
            elif progress and progress != "待組":
                state = "組立中"
            else:
                state = "未開始"

            if not progress:
                notes.append("進度空白，暫列未開始待確認")
            if occ[key] > 1:
                notes.append("同製令/料號/Type/Category重複，以出現序號保留")

            r.update({
                "編號": uid,
                "專案": raw["客戶"],
                "料號或機型": raw["P/N"],
                "名稱": raw["Type"],
                "供應商或地點": raw["組立地點"],
                "數量": 1,
                "完成數量": 1 if state == "已完成" else 0,
                "需求日期": r["客戶入庫日"],
                "預計日期": r["入庫日"],
                "狀態": state,
                "優先級": "一般",
                "負責人": raw["組立人員"],
                "客戶入庫日原文": raw["客戶入庫日"],
                "確認交期原文": txt(values[14]) if len(values) > 14 else "",
                "匯入提示": "；".join(notes),
                "來源列": line,
                "_source": "2026排程",
                "_raw": raw,
            })

            records.append(r)
            if notes:
                warnings.append(f"第 {line} 列 {raw['製令']}：" + r["匯入提示"])

        if not records:
            raise ValueError("2026排程沒有可匯入資料")

        return records, warnings
    finally:
        wb.close()


MATERIAL_FIELDS = [
    "製令","客戶","P/N","Type","Category","組立地點","組立人員",
    "組立進度","備註","發料日","入庫日","Frame/ Frame set","PU",
    "Facility","其他託外模組","RB","LP","AL","FFU","X-table","加工件","市購件","齊料"
]

def read_material(source):
    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if "在線缺料" not in wb.sheetnames:
            raise ValueError("找不到「在線缺料」工作表，請上傳倉庫物管發料原檔")

        ws = wb["在線缺料"]
        if ws.max_row > 30000:
            raise ValueError("超過 30000 列，請縮小資料範圍")

        iterator = ws.iter_rows(max_col=23, values_only=True)
        norm = lambda x: re.sub(r"\s+", "", txt(x))
        header = [norm(v) for v in next(iterator)]

        if header != [norm(f) for f in MATERIAL_FIELDS]:
            raise ValueError("在線缺料前 23 欄不符原表，請保留原欄位順序")

        rows, warnings = [], []
        seen = Counter()

        for line, values in enumerate(iterator, 2):
            if not txt(values[0]):
                continue

            r = dict(zip(MATERIAL_FIELDS, map(txt, values)))
            raw = dict(r)
            notes = []

            if isinstance(values[7], (int, float)) and 0 <= values[7] <= 1:
                r["組立進度"] = f"{values[7] * 100:g}%"

            key = json.dumps(
                [r[f] for f in ["製令", "P/N", "Type", "Category"]],
                ensure_ascii=False
            )
            seen[key] += 1

            if seen[key] > 1:
                notes.append("同組重複，以出現序號保留")

            dates = {}
            for f in ["發料日", "入庫日"]:
                dates[f], note = parse_excel_date(r[f])
                if note:
                    notes.append(f"{f}：{note}")

            ready = str(r.get("齊料", "")).strip() == "齊料"
            if not ready:
                notes.append("齊料欄未確認，列為缺料待確認")

            rid = "MAT-" + hashlib.sha256(
                (key + "#" + str(seen[key])).encode("utf-8")
            ).hexdigest()[:24]

            r.update({
                "編號": rid,
                "專案": r["客戶"],
                "料號或機型": r["P/N"],
                "名稱": r["Type"],
                "供應商或地點": r["組立地點"],
                "數量": 1,
                "完成數量": 1 if ready else 0,
                "需求日期": dates["入庫日"],
                "預計日期": "",
                "狀態": "齊料" if ready else "待確認",
                "優先級": "一般",
                "負責人": r["組立人員"],
                "匯入提示": "；".join(notes),
                "來源列": line,
                "_source": "在線缺料",
                "_raw": raw,
            })

            rows.append(r)
            if notes:
                warnings.append(f"第 {line} 列 {r['製令']}：" + r["匯入提示"])

        if not rows:
            raise ValueError("在線缺料沒有可匯入資料")

        return rows, warnings
    finally:
        wb.close()


# =========================
# 新舊 Excel 比對記憶
# =========================
def compare_schedule_row(new, old=None, stamp=None):
    r = dict(new)
    stamp = stamp or now_stamp()

    current = str(r.get("組立進度", "") or "").strip()
    previous = str((old or {}).get("組立進度", "") or "").strip()

    r["_last_imported_at"] = stamp
    r["_previous_progress"] = previous

    if not old:
        r["_progress_changed_at"] = stamp
        r["_last_compare_changed"] = None
        r["_progress_diff"] = "首次匯入／建立基準"
    elif current != previous:
        r["_progress_changed_at"] = stamp
        r["_last_compare_changed"] = True
        r["_progress_diff"] = f"{display_value(previous)} → {display_value(current)}"
    else:
        r["_progress_changed_at"] = (
            old.get("_progress_changed_at")
            or old.get("_last_imported_at")
            or stamp
        )
        r["_last_compare_changed"] = False
        r["_progress_diff"] = "無變更"

    return r

def compare_material_row(new, old=None, stamp=None):
    r = dict(new)
    stamp = stamp or now_stamp()
    changes = []
    r["_last_imported_at"] = stamp

    if old:
        for f in MATERIAL_COMPARE_FIELDS:
            ov = str(old.get(f, "") or "").strip()
            nv = str(r.get(f, "") or "").strip()
            if ov != nv:
                changes.append(f"{f}：{display_value(ov)} → {display_value(nv)}")

    if not old:
        r["_material_changed_at"] = stamp
        r["_last_compare_changed"] = None
        r["_material_diff"] = "首次匯入／建立基準"
    elif changes:
        r["_material_changed_at"] = stamp
        r["_last_compare_changed"] = True
        r["_material_diff"] = "；".join(changes)[:1800]
    else:
        r["_material_changed_at"] = (
            old.get("_material_changed_at")
            or old.get("_last_imported_at")
            or stamp
        )
        r["_last_compare_changed"] = False
        r["_material_diff"] = "無變更"

    return r

def decorate_compare(r, module):
    r = dict(r)

    if module == "production":
        days = elapsed_days(r.get("_progress_changed_at"))
        if r.get("狀態") in DONE["production"]:
            result = "已完成／停止異常追蹤"
        elif days >= 1:
            result = f"異常：進度未變動 {days} 天"
        elif r.get("_last_compare_changed") is True:
            result = "正常：進度有變更"
        elif r.get("_last_compare_changed") is False:
            result = "進度未變動（未滿 1 天）"
        elif r.get("_source") == "2026排程":
            result = "首次匯入／建立比對基準"
        else:
            result = "—"

        r["比對結果"] = result
        r["進度差異"] = r.get("_progress_diff", "—")
        r["未變動天數"] = f"{days} 天" if r.get("_source") == "2026排程" else "—"
        r["最後比對"] = str(r.get("_last_imported_at", "")).replace("T", " ") or "—"

    elif module == "material":
        days = elapsed_days(r.get("_material_changed_at"))
        if r.get("狀態") in DONE["material"]:
            result = "已完成／停止異常追蹤"
        elif days >= 1:
            result = f"異常：缺料未變動 {days} 天"
        elif r.get("_last_compare_changed") is True:
            result = "缺料有變更"
        elif r.get("_last_compare_changed") is False:
            result = "缺料未變動（未滿 1 天）"
        elif r.get("_source") == "在線缺料":
            result = "首次匯入／建立比對基準"
        else:
            result = "—"

        r["比對結果"] = result
        r["缺料差異"] = r.get("_material_diff", "—")
        r["未變動天數"] = f"{days} 天" if r.get("_source") == "在線缺料" else "—"
        r["最後比對"] = str(r.get("_last_imported_at", "")).replace("T", " ") or "—"

    return r


# =========================
# 風險判斷
# =========================
def risk(row, module):
    r = row

    if r.get("狀態") in DONE[module]:
        return "🟢", "已完成／不占用"

    if module == "warehouse":
        return "🟢", "使用中"

    if module == "production":
        due = parse_day(r.get("入庫日") or r.get("預計日期"))
        stagnant = elapsed_days(r.get("_progress_changed_at")) if r.get("_source") == "2026排程" else 0

        if due and due < today():
            overdue = (today() - due).days
            suffix = f"／進度未變 {stagnant} 天" if stagnant >= 1 else ""
            return "🔴", f"入庫逾期 {overdue} 天{suffix}"

        if stagnant >= 1:
            return "🔴", f"進度異常／未變動 {stagnant} 天"

        if not due:
            return "🟡", "入庫日未填"

        if r.get("優先級") == "急件" or due <= today() + timedelta(days=3):
            return "🟡", "入庫日三日內／關注"

        return "🟢", "正常"

    if module == "material" and r.get("_source") == "在線缺料":
        if r.get("齊料") == "齊料":
            return "🟢", "齊料"

        due = parse_day(r.get("入庫日") or r.get("需求日期"))
        stagnant = elapsed_days(r.get("_material_changed_at"))

        if due and due < today():
            overdue = (today() - due).days
            suffix = f"／缺料未變 {stagnant} 天" if stagnant >= 1 else ""
            return "🔴", f"缺料逾期 {overdue} 天{suffix}"

        if stagnant >= 1:
            return "🔴", f"缺料異常／未變動 {stagnant} 天"

        if not due:
            return "🟡", "入庫日未填／缺料待確認"

        return "🟡", "缺料待處理"

    due = parse_day(r.get("需求日期"))
    eta = parse_day(r.get("預計日期"))

    if r.get("狀態") == "不合格" or (due and (due < today() or (eta and eta > due))):
        return "🔴", "逾期／異常"

    if r.get("優先級") == "急件" or (due and due <= today() + timedelta(days=3)):
        return "🟡", "急件／三日內到期"

    if not due:
        return "🟡", "需求日期未填"

    return "🟢", "正常"


def make_table(rows, module):
    out = []
    for row in rows:
        r = decorate_compare(row, module)
        lamp, label = risk(r, module)

        # 隱藏內部欄位，保留 Excel 實際欄位與戰情欄位
        public = {k: v for k, v in r.items() if not k.startswith("_")}
        out.append({
            "風險": lamp,
            "風險說明": label,
            **public
        })

    return pd.DataFrame(out).fillna("")


# =========================
# Secrets / 登入
# =========================
st.title("超慧科技｜生管資材戰情看板")
st.caption("Excel 為主資料來源；雲端資料庫只保存操作記憶、比對基準、IQC/倉庫資料與歷史狀態。")

try:
    DATABASE_URL = str(st.secrets["DATABASE_URL"]).strip()
    ADMIN_PASSWORD = str(st.secrets["admin_password"])
    VIEWER_PASSWORD = str(st.secrets["viewer_password"])
except Exception:
    st.error("首次部署尚未完成：請在 Streamlit → Settings → Secrets 設定 DATABASE_URL、admin_password、viewer_password。")
    st.code(
        'DATABASE_URL = "postgresql://USER:PASSWORD@HOST:5432/DATABASE?sslmode=require"\n'
        'admin_password = "請設定至少12碼管理者密碼"\n'
        'viewer_password = "請設定另一組至少12碼瀏覽者密碼"',
        language="toml"
    )
    st.info("注意：Excel 仍是主資料來源。DATABASE_URL 只用來確保重新部署、休眠、換電腦後，操作記憶不會歸零。")
    st.stop()

if (
    len(ADMIN_PASSWORD) < 12
    or len(VIEWER_PASSWORD) < 12
    or ADMIN_PASSWORD == VIEWER_PASSWORD
):
    st.error("admin_password 與 viewer_password 必須不同，且都至少 12 碼。")
    st.stop()

try:
    db_initialize(DATABASE_URL)
except Exception as e:
    st.error("無法連線到持久記憶資料庫。為避免操作記憶遺失，系統已停止寫入。")
    st.caption(f"技術訊息：{type(e).__name__}")
    st.stop()

if not st.session_state.get("role"):
    with st.form("login_form"):
        role = st.selectbox(
            "登入身分",
            ["viewer", "admin"],
            format_func=lambda x: "管理者" if x == "admin" else "瀏覽者"
        )
        password = st.text_input("密碼", type="password")

        if st.form_submit_button("登入", use_container_width=True):
            target = ADMIN_PASSWORD if role == "admin" else VIEWER_PASSWORD
            if hmac.compare_digest(password.encode(), target.encode()):
                st.session_state.role = role
                st.rerun()
            else:
                st.error("密碼不正確")
    st.stop()

admin = st.session_state.role == "admin"

try:
    all_data = db_load(DATABASE_URL)
except Exception as e:
    st.error("讀取操作記憶失敗，為避免顯示空白假資料，系統已停止。")
    st.caption(f"技術訊息：{type(e).__name__}")
    st.stop()


# =========================
# 側邊選單
# =========================
with st.sidebar:
    st.markdown("## 📊 生管資材")
    st.success("Excel 主資料＋持久操作記憶")
    page = st.radio(
        "功能",
        ["戰情總覽", *MODULES.values(), "操作記憶／備份"]
    )

    if st.button("🔄 重新整理", use_container_width=True):
        st.rerun()

    if st.button("🚪 登出", use_container_width=True):
        st.session_state.clear()
        st.rerun()


# =========================
# 戰情總覽
# =========================
if page == "戰情總覽":
    cols = st.columns(4)

    for col, (module, title) in zip(cols, MODULES.items()):
        rows = [r for r, _ in all_data[module].values()]
        unfinished = sum(1 for r in rows if r.get("狀態") not in DONE[module])
        col.metric(title, len(rows), f"未結案 {unfinished}")

    st.divider()

    for module, title in MODULES.items():
        st.subheader(title)
        rows = [r for r, _ in all_data[module].values()]

        if module == "production" and rows:
            counts = Counter(
                str(r.get("組立進度") or r.get("狀態") or "未分類")
                for r in rows
            )
            st.bar_chart(pd.Series(counts, name="台數"))

        if rows:
            st.dataframe(
                make_table(rows, module),
                use_container_width=True,
                hide_index=True,
            )
        else:
            st.info("目前尚無資料。")


# =========================
# 操作記憶 / 備份
# =========================
elif page == "操作記憶／備份":
    if not admin:
        st.info("此頁僅管理者可使用。")
        st.stop()

    st.subheader("操作記憶狀態")

    meta = db_get_all_meta(DATABASE_URL)
    if meta:
        meta_rows = []
        for k, item in meta.items():
            meta_rows.append({
                "項目": k,
                "內容": json.dumps(item["value"], ensure_ascii=False)
                        if not isinstance(item["value"], str) else item["value"],
                "更新時間": item["updated_at"],
            })
        st.dataframe(pd.DataFrame(meta_rows), use_container_width=True, hide_index=True)
    else:
        st.info("尚無匯入紀錄。")

    backup = {
        "exported_at": now_stamp(),
        "modules": {
            m: [r for r, _ in entries.values()]
            for m, entries in all_data.items()
        },
        "meta": meta,
    }

    st.download_button(
        "⬇️ 下載全部操作記憶 JSON 備份",
        data=json.dumps(backup, ensure_ascii=False, indent=2),
        file_name=f"dashboard_backup_{today().isoformat()}.json",
        mime="application/json",
        use_container_width=True,
    )

    st.warning("更新 GitHub / Streamlit 時，請保留同一組 DATABASE_URL。更換成新的空白資料庫會看不到原操作記憶。")


# =========================
# 各模組
# =========================
else:
    module = next(m for m, title in MODULES.items() if title == page)
    entries = all_data[module]
    rows = [r for r, _ in entries.values()]

    main, right = st.columns([3, 1])

    with main:
        keyword = st.text_input("搜尋製令／料號／專案／備註")

        if keyword:
            selected = [
                r for r in rows
                if keyword.lower() in json.dumps(r, ensure_ascii=False).lower()
            ]
        else:
            selected = rows

        frame = make_table(selected, module)

        st.dataframe(
            frame,
            use_container_width=True,
            hide_index=True,
        )

        st.download_button(
            "下載目前清單 CSV",
            data=frame.to_csv(index=False).encode("utf-8-sig"),
            file_name=f"{module}_{today().isoformat()}.csv",
            mime="text/csv",
        )

    with right:
        st.subheader("Excel 上傳")

        if admin and module in ["production", "material"]:
            upload_label = (
                "上傳 2026排程 Excel"
                if module == "production"
                else "上傳倉庫物管發料 Excel"
            )

            f = st.file_uploader(
                upload_label,
                type=["xlsx"],
                key=f"upload_{module}",
            )

            if f and st.button("① 預覽匯入", use_container_width=True):
                try:
                    parser = read_schedule if module == "production" else read_material
                    imported, warnings = parser(io.BytesIO(f.getvalue()))
                    st.session_state[f"preview_{module}"] = {
                        "rows": imported,
                        "warnings": warnings,
                    }
                except Exception as e:
                    st.error(str(e))

            preview = st.session_state.get(f"preview_{module}")

            if preview:
                imported = preview["rows"]
                warnings = preview["warnings"]

                st.success(f"待匯入 {len(imported)} 筆")
                st.caption("同編號會更新；Excel 未包含的舊資料不刪除；人工備註與優先級保留。")

                st.dataframe(
                    make_table(imported[:100], module),
                    use_container_width=True,
                    hide_index=True,
                    height=360,
                )

                if len(imported) > 100:
                    st.caption("預覽僅顯示前 100 筆，實際會全部匯入。")

                with st.expander(f"匯入提示 ({len(warnings)})"):
                    if warnings:
                        for w in warnings[:200]:
                            st.write("• " + w)
                    else:
                        st.write("無")

                if st.button("② 確認寫入並保存記憶", type="primary", use_container_width=True):
                    try:
                        saved = db_save_import(
                            DATABASE_URL,
                            module,
                            imported,
                            entries
                        )
                        st.session_state.pop(f"preview_{module}", None)
                        st.success(f"完成：{saved} 筆已寫入，舊操作記憶已保留。")
                        st.rerun()
                    except Exception as e:
                        st.error(f"寫入失敗：{e}")

        elif not admin:
            st.info("瀏覽者只能查看。")
        else:
            st.caption("IQC 與倉庫資料由下方表單維護。")

    # 管理者人工操作
    if admin:
        st.divider()
        st.subheader("新增／編輯資料")

        rid = st.selectbox("選擇資料", ["新增", *list(entries.keys())])

        if rid == "新增":
            old = {}
            version = None
        else:
            old, version = entries[rid]

        with st.form(f"edit_{module}_{rid}"):
            values = {}
            cols = st.columns(3)

            for i, field in enumerate(FIELDS):
                with cols[i % 3]:
                    if field == "狀態":
                        opts = STATES[module]
                        current = old.get(field)
                        idx = opts.index(current) if current in opts else 0
                        values[field] = st.selectbox(field, opts, index=idx)

                    elif field == "優先級":
                        opts = ["一般", "急件"]
                        current = old.get(field)
                        idx = opts.index(current) if current in opts else 0
                        values[field] = st.selectbox(field, opts, index=idx)

                    else:
                        disabled = field == "編號" and rid != "新增"
                        values[field] = st.text_input(
                            field,
                            value=str(old.get(field, "")),
                            disabled=disabled,
                        )

            submitted = st.form_submit_button("💾 儲存操作記憶", use_container_width=True)

            if submitted:
                try:
                    row = dict(old)
                    row.update(values)

                    if module == "iqc" and not row.get("收料時間"):
                        row["收料時間"] = now_stamp()

                    # 人工修改備註時留下時間記憶
                    if str(row.get("備註", "")) != str(old.get("備註", "")):
                        row["_manual_note_updated_at"] = now_stamp()

                    db_save_manual(
                        DATABASE_URL,
                        module,
                        row,
                        expected_version=version
                    )
                    st.success("已保存，重新部署後仍會保留。")
                    st.rerun()

                except Exception as e:
                    st.error(str(e))
