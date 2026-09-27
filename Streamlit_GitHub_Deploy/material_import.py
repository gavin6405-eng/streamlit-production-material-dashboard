import hashlib,json,re
from collections import Counter
from openpyxl import load_workbook
from schedule_import import txt,dt
FIELDS=['製令','客戶','P/N','Type','Category','組立地點','組立人員','組立進度','備註','發料日','入庫日','Frame/ Frame set','PU','Facility','其他託外模組','RB','LP','AL','FFU','X-table','加工件','市購件','齊料']
def read_material(source):
    wb=load_workbook(source,read_only=True,data_only=True)
    try:
        if '在線缺料' not in wb.sheetnames:raise ValueError('找不到「在線缺料」工作表，請上傳倉庫物管發料原檔')
        ws=wb['在線缺料']
        if ws.max_row>20000:raise ValueError('超過20000列，請縮小範圍')
        it=ws.iter_rows(max_col=23,values_only=True)
        norm=lambda x:re.sub(r'\s+','',txt(x))
        if [norm(v) for v in next(it)]!=[norm(f) for f in FIELDS]:raise ValueError('在線缺料前23欄不符原表，請保留原欄位順序')
        rows=[];warnings=[];seen=Counter()
        for line,v in enumerate(it,2):
            if not txt(v[0]):continue
            r=dict(zip(FIELDS,map(txt,v)));raw=dict(r);notes=[]
            if isinstance(v[7],(int,float)) and 0<=v[7]<=1:r['組立進度']=f'{v[7]*100:g}%'
            key=json.dumps([r[f] for f in ['製令','P/N','Type','Category']],ensure_ascii=False);seen[key]+=1
            if seen[key]>1:notes.append('同組重複，以出現序號保留')
            dates={}
            for f in ['發料日','入庫日']:
                dates[f],note=dt(r[f])
                if note:notes.append(f+'：'+note)
            ready=r['齊料']=='齊料'
            if not ready:notes.append('齊料欄未確認，請依原文判讀，非缺料數量')
            r.update({'編號':'MAT-'+hashlib.sha256((key+'#'+str(seen[key])).encode()).hexdigest()[:24], '專案':r['客戶'],'料號或機型':r['P/N'],'名稱':r['Type'],'供應商或地點':r['組立地點'],'數量':1,'完成數量':int(ready),'需求日期':dates['入庫日'],'預計日期':'','狀態':'齊料' if ready else '待確認','優先級':'一般','負責人':r['組立人員'],'匯入提示':'；'.join(notes),'_source':'在線缺料','_raw':raw,'來源列':line})
            rows.append(r)
            if notes:warnings.append(f'第{line}列 {r["製令"]}：'+r['匯入提示'])
        if not rows:raise ValueError('在線缺料沒有製令資料')
        return rows,warnings
    finally:wb.close()
