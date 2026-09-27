import json, hashlib, re
from collections import Counter
from datetime import date, datetime
from openpyxl import load_workbook
FIELDS=['製令','客戶','P/N','Type','Category','組立地點','組立人員','組立進度','備註','發料日','入庫日','保稅','客戶入庫日']
def txt(v):
    if v is None:return ''
    if isinstance(v,(date,datetime)):return v.strftime('%Y-%m-%d')
    return str(v).strip()
def dt(v):
    s=txt(v)
    if s in ['', '-', '--','TBD','託']:return '', ''
    for fmt in ['%Y-%m-%d','%Y/%m/%d','%Y%m%d','%Y%m/%d','%Y/%m%d']:
        try:return datetime.strptime(s.splitlines()[-1].strip(),fmt).date().isoformat(),('多行日期採最後一行' if '\n' in s else '')
        except ValueError:pass
    return '', '日期無法辨識，請確認原文'
def read_schedule(source):
    wb=load_workbook(source,read_only=True,data_only=True)
    try:
        if '2026排程' not in wb.sheetnames:raise ValueError('找不到「2026排程」工作表')
        ws=wb['2026排程']
        if ws.max_row>20000:raise ValueError('超過20000列')
        iterator=ws.iter_rows(max_col=15,values_only=True)
        header=[re.sub(r'\s+','',txt(v)) for v in next(iterator)]
        if header[:13]!=FIELDS:raise ValueError('前13欄與原始2026排程格式不符，請保留欄位順序')
        records=[];occ=Counter();warnings=[]
        for line,v in enumerate(iterator,2):
            if not txt(v[0]):continue
            raw=dict(zip(FIELDS,map(txt,v[:13])));r=dict(raw);notes=[]
            key=json.dumps([raw[k] for k in ['製令','P/N','Type','Category']],ensure_ascii=False);occ[key]+=1
            uid='SCH26-'+hashlib.sha256((key+'#'+str(occ[key])).encode()).hexdigest()[:24]
            for f in ['發料日','入庫日','客戶入庫日']:
                r[f],note=dt(raw[f])
                if note:notes.append(f+'：'+note)
            progress=r['組立進度']
            state='已完成' if progress in ['已完工','已完成'] else '暫停' if '暫停' in progress else '檢驗中' if progress in ['Q','待Q','R','Q / R'] else '組立中' if progress and progress!='待組' else '未開始'
            if not progress:notes.append('進度空白，暫列未開始待確認')
            if occ[key]>1:notes.append('同製令/料號/Type/Category重複，以出現序號保留')
            r.update({'編號':uid,'專案':raw['客戶'],'料號或機型':raw['P/N'],'名稱':raw['Type'],'供應商或地點':raw['組立地點'],'數量':1,'完成數量':int(state=='已完成'),'需求日期':r['客戶入庫日'],'預計日期':r['入庫日'],'狀態':state,'優先級':'一般','負責人':raw['組立人員'],'客戶入庫日原文':raw['客戶入庫日'],'確認交期原文':txt(v[14]),'匯入提示':'；'.join(notes),'來源列':line,'_source':'2026排程','_raw':raw})
            records.append(r)
            if notes:warnings.append(f'第{line}列 {raw["製令"]}：'+r['匯入提示'])
        if not records:raise ValueError('工作表沒有製令資料')
        return records,warnings
    finally:wb.close()
